from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session
from sqlalchemy import desc
from app.services.graph_service import get_graph_service, get_available_tenants
from app.services.ad_service import get_all_users_last_logon
from app.core.database import get_db
from app.models.license_audit import LicenseAudit
import datetime
import os
from pydantic import BaseModel

router = APIRouter()

class UserNote(BaseModel):
    userId: str
    userPrincipalName: str = ""
    displayName: str = ""
    note: str
    action: str = "NOTE" # "NOTE" or "RECOVERED"
    licenses: str = ""
    savedUsd: int = 0

class ExportPayload(BaseModel):
    sheetName: str
    data: list

@router.get("/tenants")
async def get_tenants():
    return {"success": True, "data": get_available_tenants()}


# ── Caché global para licencias inactivas ─────────────────────────────────────
# Keyed by f"{tenant}:{days}" para soportar distintos filtros
_inactive_cache: dict = {}
_INACTIVE_CACHE_TTL = 600  # 10 minutos

# Estado de carga activa por clave de caché
_inactive_loading: dict = {}  # key -> {"running": bool, "progress": int, "total": int}

@router.get("/inactive/status")
async def get_inactive_licenses_status(days: int = 90, tenant: str = "1"):
    """Retorna el estado de la carga de licencias inactivas (para mostrar progreso al frontend)."""
    cache_key = f"{tenant}:{days}"
    cached = _inactive_cache.get(cache_key)
    loading = _inactive_loading.get(cache_key, {})
    
    if cached and (datetime.datetime.now().timestamp() - cached["ts"]) < _INACTIVE_CACHE_TTL:
        return {
            "ready": True,
            "cached": True,
            "count": len(cached["data"]),
            "progress": len(cached["data"]),
            "total": len(cached["data"])
        }
    
    return {
        "ready": False,
        "cached": False,
        "running": loading.get("running", False),
        "progress": loading.get("progress", 0),
        "total": loading.get("total", 0)
    }

@router.get("/inactive")
async def get_inactive_licenses(days: int = 90, tenant: str = "1", db: Session = Depends(get_db)):
    """
    Retorna los usuarios que tienen licencias asignadas 
    pero que no han iniciado sesión en los últimos N días.
    Usa caché en memoria (TTL 10 min) para respuesta instantánea en la segunda visita.
    """
    cache_key = f"{tenant}:{days}"
    
    # Servir desde caché si está fresco
    cached = _inactive_cache.get(cache_key)
    if cached and (datetime.datetime.now().timestamp() - cached["ts"]) < _INACTIVE_CACHE_TTL:
        # Inyectar notas frescas de DB (esto es rápido)
        all_notes = db.query(LicenseAudit).filter(LicenseAudit.action == "NOTE").all()
        notes_db = {note.user_id: note.note for note in all_notes}
        results = [dict(u, note=notes_db.get(u.get("id"), "")) for u in cached["data"]]
        return {"success": True, "data": results, "from_cache": True}
    
    try:
        cache_key_loading = cache_key
        _inactive_loading[cache_key_loading] = {"running": True, "progress": 0, "total": 0}
        
        svc = get_graph_service(tenant)
        results = await svc.get_inactive_licensed_users(days_threshold=days)
        
        _inactive_loading[cache_key_loading] = {"running": True, "progress": len(results), "total": len(results)}
        
        # Load notes from DB
        all_notes = db.query(LicenseAudit).filter(LicenseAudit.action == "NOTE").all()
        notes_db = {note.user_id: note.note for note in all_notes}
        
        try:
            ad_logons_dict = get_all_users_last_logon()
            now = datetime.datetime.now(datetime.timezone.utc)
            
            for user in results:
                immutable_id = user.get("onPremisesImmutableId")
                raw_username = user.get("onPremisesSamAccountName")
                username = raw_username.lower() if raw_username else ""
                
                local_logon_iso = None
                if immutable_id and immutable_id in ad_logons_dict["by_guid"]:
                    local_logon_iso = ad_logons_dict["by_guid"][immutable_id]
                elif username and username in ad_logons_dict["by_sam"]:
                    local_logon_iso = ad_logons_dict["by_sam"][username]
                
                user["localAdLastLogon"] = local_logon_iso
                user["daysInactiveLocalAd"] = 9999
                
                if local_logon_iso:
                    try:
                        dt_str = local_logon_iso.replace('Z', '+00:00')
                        dt = datetime.datetime.fromisoformat(dt_str)
                        if dt.tzinfo is None:
                            dt = dt.replace(tzinfo=datetime.timezone.utc)
                        user["daysInactiveLocalAd"] = max(0, (now - dt).days)
                    except Exception:
                        pass
                        
                user["note"] = notes_db.get(user.get("id"), "")
                
        except Exception as ad_err:
            print(f"Error fetching AD logons: {ad_err}")

        # Guardar en caché
        _inactive_cache[cache_key] = {"ts": datetime.datetime.now().timestamp(), "data": results}
        _inactive_loading[cache_key_loading] = {"running": False, "progress": len(results), "total": len(results)}
            
        return {"success": True, "data": results, "from_cache": False}
    except Exception as e:
        _inactive_loading[cache_key] = {"running": False, "progress": 0, "total": 0}
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/notes")
async def save_user_note(payload: UserNote, db: Session = Depends(get_db)):
    """Guarda o actualiza una nota de auditoría para un usuario, o registra una licencia recuperada."""
    try:
        if payload.action == "NOTE":
            # Actualizar o crear nota
            existing = db.query(LicenseAudit).filter(
                LicenseAudit.user_id == payload.userId,
                LicenseAudit.action == "NOTE"
            ).first()
            if existing:
                existing.note = payload.note
                existing.timestamp = datetime.datetime.utcnow()
            else:
                new_note = LicenseAudit(
                    user_id=payload.userId,
                    user_principal_name=payload.userPrincipalName,
                    display_name=payload.displayName,
                    action="NOTE",
                    note=payload.note
                )
                db.add(new_note)
        elif payload.action == "RECOVERED":
            # Verificar con Microsoft Graph si realmente se quitó la licencia (usa Tenant 1 por defecto para recoveries o asume que el token no es crítico)
            # Idealmente, recibiríamos el tenant. Por ahora se usa el 1.
            svc = get_graph_service("1")
            current_licenses = await svc.get_user_licenses(payload.userId)
            if len(current_licenses) > 0:
                raise ValueError("Esta cuenta aún cuenta con licencia activa. Retírala primero en Microsoft 365 para poder registrar la recuperación.")

            # Crear registro inmutable de recuperación
            new_audit = LicenseAudit(
                user_id=payload.userId,
                user_principal_name=payload.userPrincipalName,
                display_name=payload.displayName,
                action="RECOVERED",
                note=payload.note,
                licenses=payload.licenses,
                saved_usd=payload.savedUsd
            )
            db.add(new_audit)
            
        db.commit()
        return {"success": True, "message": "Acción registrada correctamente"}
    except ValueError as ve:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(ve))
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/audit/history")
def get_audit_history(db: Session = Depends(get_db)):
    """Obtiene el histórico de licencias recuperadas."""
    try:
        records = db.query(LicenseAudit).filter(LicenseAudit.action == "RECOVERED").order_by(desc(LicenseAudit.timestamp)).all()
        return {"success": True, "data": records}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/skus")
async def get_skus(tenant: str = "1"):
    """
    Retorna las cuotas reales de licencias (Subscribed SKUs) del Tenant.
    """
    try:
        svc = get_graph_service(tenant)
        results = await svc.get_subscribed_skus()
        return {"success": True, "data": results}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/export")
async def export_to_excel(payload: ExportPayload):
    """
    Exporta la tabla a un archivo Excel en la carpeta de Documentos del usuario.
    Maneja múltiples hojas según el filtro enviado.
    """
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
        
        # Ruta en Mis Documentos del usuario actual del sistema Windows
        docs_folder = os.path.join(os.path.expanduser("~"), "Documents")
        file_path = os.path.join(docs_folder, "Limpieza_Licencias.xlsx")
        
        wb = None
        if os.path.exists(file_path):
            try:
                wb = openpyxl.load_workbook(file_path)
            except Exception:
                # Si falla abrirlo por alguna razón, creamos uno nuevo
                pass
                
        if not wb:
            wb = openpyxl.Workbook()
            # Eliminar la hoja por defecto "Sheet"
            if "Sheet" in wb.sheetnames:
                del wb["Sheet"]
                
        # Limpiar nombre de hoja (max 31 chars, no invalid chars)
        safe_sheet_name = "".join([c for c in payload.sheetName if c not in r'[]:*?/\\']).strip()
        safe_sheet_name = safe_sheet_name[:31] if safe_sheet_name else "Export"
        
        if safe_sheet_name in wb.sheetnames:
            del wb[safe_sheet_name]
            
        ws = wb.create_sheet(title=safe_sheet_name)
        
        # Escribir encabezados
        headers = [
            "Identidad", "Usuario (UPN)", "Cuenta AD", "Estado", 
            "Licencias Asignadas", "Inicio AD (Local)", "Inicio M365 (Interactivo)", 
            "Inicio M365 (Fondo)", "Días Inactivos (Mínimo)", "Nota / Auditoría"
        ]
        
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = openpyxl.styles.Font(bold=True)
            
        # Escribir datos
        for row_num, user in enumerate(payload.data, 2):
            # Parse licenses
            licenses_str = ", ".join(user.get("licenses", [])) if user.get("licenses") else "Sin licencias"
            
            # Format dates (just keeping them as strings or empty if null)
            ad_local = user.get("localAdLastLogon") or "Nunca"
            m365_inter = user.get("lastSignInDateTimeAd") or "Nunca"
            m365_fondo = user.get("lastSignInDateTimeOutlook") or "Nunca"
            
            # Days calculation
            ad_d = user.get("daysInactiveLocalAd", 9999)
            m365_d = user.get("daysInactiveAd", 9999)
            fondo_d = user.get("daysInactiveOutlook", 9999)
            min_days = min(ad_d, m365_d, fondo_d)
            if min_days == 9999:
                min_days = "Nunca ha iniciado"
                
            estado = "Habilitado" if user.get("accountEnabled") else "Inhabilitado"
            if user.get("isBlocked"):
                estado = "Bloqueado"
                
            ws.cell(row=row_num, column=1, value=user.get("displayName", ""))
            ws.cell(row=row_num, column=2, value=user.get("userPrincipalName", ""))
            ws.cell(row=row_num, column=3, value=user.get("onPremisesSamAccountName", ""))
            ws.cell(row=row_num, column=4, value=estado)
            ws.cell(row=row_num, column=5, value=licenses_str)
            ws.cell(row=row_num, column=6, value=ad_local)
            ws.cell(row=row_num, column=7, value=m365_inter)
            ws.cell(row=row_num, column=8, value=m365_fondo)
            ws.cell(row=row_num, column=9, value=str(min_days))
            ws.cell(row=row_num, column=10, value=user.get("note", ""))
            
        # Auto-ajustar ancho de columnas
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = min(adjusted_width, 50) # Cap at 50

        wb.save(file_path)
        
        from fastapi.responses import FileResponse
        return FileResponse(
            path=file_path, 
            filename="Limpieza_Licencias.xlsx", 
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        
    except PermissionError:
        raise HTTPException(status_code=500, detail="Error de permisos. ¿Tienes el archivo de Excel abierto? Ciérralo y vuelve a intentar.")
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))
